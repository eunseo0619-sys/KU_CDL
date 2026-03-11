import os
import re
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import numpy as np

DOCS_DIR = os.path.join(os.path.dirname(__file__), "docs")

DOC_FILES = {
    "인수인계": "CDL 1F, 2F 인수인계_2026.02.06 (1).docx",
    "도서관 규정": "고려대 도서관 규정 정리 - cdl 관련.docx",
}


def extract_chunks(path: str) -> list[dict]:
    """워드 파일에서 의미 단위 청크를 추출합니다."""
    doc = Document(path)
    chunks = []
    current_heading = ""
    para_buffer = []

    def flush_buffer():
        if para_buffer:
            text = "\n".join(para_buffer)
            if len(text.strip()) >= 15:
                chunks.append({"text": text, "heading": current_heading})
            para_buffer.clear()

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if para.style.name.startswith("Heading"):
            flush_buffer()
            current_heading = text
        else:
            para_buffer.append(text)
            # 버퍼가 충분히 쌓이면 청크로 저장
            if len("\n".join(para_buffer)) >= 80:
                flush_buffer()

    flush_buffer()

    # 표는 전체를 하나의 청크로
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            table_text = "\n".join(rows)
            if len(table_text) >= 15:
                chunks.append({"text": table_text, "heading": current_heading})

    return chunks


def load_documents() -> dict:
    """두 워드 파일을 읽어 검색 가능한 형태로 반환합니다."""
    all_chunks = []
    missing = []

    for label, filename in DOC_FILES.items():
        filepath = os.path.join(DOCS_DIR, filename)
        if not os.path.exists(filepath):
            missing.append(filename)
            continue
        for chunk in extract_chunks(filepath):
            all_chunks.append({
                "source": label,
                "heading": chunk["heading"],
                "text": chunk["text"],
            })

    if not all_chunks:
        return {"chunks": [], "vectorizer": None, "matrix": None, "missing": missing}

    texts = [c["text"] for c in all_chunks]
    vectorizer = TfidfVectorizer(
        analyzer="char_wb",
        ngram_range=(2, 4),
        max_features=20000,
    )
    matrix = vectorizer.fit_transform(texts)

    return {
        "chunks": all_chunks,
        "vectorizer": vectorizer,
        "matrix": matrix,
        "missing": missing,
    }


def search_documents(docs: dict, query: str, top_k: int = 5) -> list[dict]:
    """질문과 가장 관련 있는 청크를 반환합니다."""
    if not docs["chunks"]:
        return []

    q_vec = docs["vectorizer"].transform([query])
    scores = cosine_similarity(q_vec, docs["matrix"]).flatten()
    top_indices = np.argsort(scores)[::-1][:top_k]

    results = []
    for idx in top_indices:
        if scores[idx] < 0.05:
            continue
        chunk = docs["chunks"][idx]
        results.append({
            "source": chunk["source"],
            "heading": chunk["heading"],
            "text": chunk["text"],
            "score": float(scores[idx]),
        })

    return results
